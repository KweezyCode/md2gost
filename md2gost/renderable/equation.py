from typing import Generator
from uuid import uuid4

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import CT_Tbl
from docx.shared import Pt, Twips
from docx.table import Table

from .caption import CaptionInfo
from .requires_numbering import RequiresNumbering
from ..layout_tracker import LayoutState
from ..renderable import Renderable
from ..rendered_info import RenderedInfo
from ..util import create_element
from ..latex_math import latex_to_omml


_HEIGHT = Pt(50)


class Equation(Renderable, RequiresNumbering):
    def __init__(self, parent, latex_formula: str, caption_info: CaptionInfo):
        super().__init__("Формула", caption_info.unique_name if caption_info else None)
        word_math = latex_to_omml(latex_formula)

        sect = parent.part.document.sections[0]

        # todo: style inheritance
        left_margin = Twips(int(parent.part.styles["Normal Table"]._element.xpath("w:tblPr/w:tblCellMar/w:left")[0].attrib["{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w"]))
        right_margin = Twips(int(parent.part.styles["Normal Table"]._element.xpath("w:tblPr/w:tblCellMar/w:right")[0].attrib["{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w"]))

        table_width = sect.page_width - sect.right_margin - sect.left_margin + left_margin + right_margin

        self._table = table = Table(CT_Tbl.new_tbl(1, 2, table_width), parent)

        left_cell = table.cell(0,0)
        right_cell = table.cell(0,1)
        right_cell.width = Pt(30)
        left_cell.width = table_width - right_cell.width

        table.rows[0].height = _HEIGHT  # TODO: implement proper size

        left_paragraph = left_cell.paragraphs[0]
        left_paragraph.style = "Formula Content"
        left_paragraph._p.append(word_math)
        left_cell.vertical_alignment = \
            WD_CELL_VERTICAL_ALIGNMENT.CENTER

        uid = uuid4().hex

        right_paragraph = right_cell.paragraphs[0]
        right_paragraph.style = "Formula Numbering"
        right_paragraph._p.append(create_element("w:r", "("))
        if caption_info and caption_info.unique_name:
            right_paragraph._p.append(create_element("w:bookmarkStart", {
                "w:id": uid,
                "w:name": caption_info.unique_name
            }))
        self._numbering_run = create_element("w:r", "?")
        right_paragraph._p.append(
            create_element("w:fldSimple", {
                "w:instr": f"SEQ Формула \\* ARABIC"
            }, [self._numbering_run]))
        if caption_info and caption_info.unique_name:
            right_paragraph._p.append(create_element("w:bookmarkEnd", {
                "w:id": uid
            }))
        right_paragraph._p.append(create_element("w:r", ")"))
        right_cell.vertical_alignment = \
            WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def set_number(self, number: int):
        self._numbering_run.text = str(number)

    def render(self, previous_rendered: RenderedInfo, layout_state: LayoutState) -> Generator[
            "RenderedInfo | Renderable", None, None]:
        height = _HEIGHT

        if height > layout_state.remaining_page_height:
            height += layout_state.remaining_page_height

        yield RenderedInfo(self._table, height)
